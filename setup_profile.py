import os
import csv
import shutil
import subprocess
import time
import calendar
from datetime import datetime
import requests
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font
from playwright.sync_api import sync_playwright
from azure.communication.email import EmailClient
from dotenv import load_dotenv

load_dotenv()

# =========================
# CONSTANTS
# =========================
WIS_URL = "https://pwc.moveinsync.com/WP/employee.jsp#WorkInSyncDashboard"
EDGE_EXE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
CDP_PORT = 9222
TIMEOUT_MS = 30000
SEARCH_WAIT_SEC = 2

INPUT_FILE = "employee_input.xlsx"
OUTPUT_FILE = "attendance_output.xlsx"
TRACKER_FILE = "Tracker.xlsx"
TEST_TRACKER_FILE = "Test_Tracker.xlsx"
FAILED_AWARENESS_CSV = "failed_awareness_mails.csv"

ABHINANDAN_ID = "101675341"
ABHINANDAN_NAME = "Abhinandan Roy"
ABHINANDAN_EMAIL = "abhinandan.roy@pwc.com"
DEBORSHI_EMAIL = "deborshi.som@pwc.com"
SHREYASI_EMAIL = "shreyasi.dutta@pwc.com"

ACS_CONNECTION = os.getenv("ACS_CONNECTION")
ACS_SENDER_EMAIL = "DoNotReply@fdd1da09-f617-4c54-b81f-143d8de4b93e.azurecomm.net"

ONEDRIVE_FULL_REPORT_URL = os.getenv(
    "ONEDRIVE_FULL_REPORT_URL",
    "https://pwcindia-my.sharepoint.com/:u:/r/personal/abhinandan_roy_pwc_com/Documents/work-in-sync/exported_reports/month_end_full_report_latest.html?d=w249b7a4d1d7d4e478c3bec391032ad2e&csf=1&web=1&e=9RyGg9"
).strip()

if not ACS_CONNECTION:
    raise ValueError("ACS_CONNECTION environment variable is not set. Please add it in .env file.")

MONTHLY_TARGET_DAYS = 10
MAIL_RETRY_COUNT = 6
MAIL_THROTTLE_DELAY_SEC = 20

DEBUG_DIR = "debug_output"
SCREENSHOT_DIR = os.path.join(DEBUG_DIR, "screenshots")
HTML_DIR = os.path.join(DEBUG_DIR, "html")
TEXT_DIR = os.path.join(DEBUG_DIR, "text")
REPORT_EXPORT_DIR = "exported_reports"
BACKUP_DIR = "excel_backups"
EDGE_PROFILE_DIR = os.path.join(os.getcwd(), "edge_debug_profile")

os.makedirs(SCREENSHOT_DIR, exist_ok=True)
os.makedirs(HTML_DIR, exist_ok=True)
os.makedirs(TEXT_DIR, exist_ok=True)
os.makedirs(REPORT_EXPORT_DIR, exist_ok=True)
os.makedirs(BACKUP_DIR, exist_ok=True)
os.makedirs(EDGE_PROFILE_DIR, exist_ok=True)

# =========================
# XPATHS / SELECTORS
# =========================
VIEW_TEAM_CALENDAR_XPATH = "xpath=//*[@class='mis-btn mis-none mis-btn-md' and text()='View Team Calendar']"
SEARCH_XPATH = 'xpath=//*[@placeholder="Search by name, ID or email "]'
CLEAR_ALL_XPATH = "xpath=//*[text()='Clear All']"

ABHINANDAN_STATUS_XPATH = "xpath=(//div[@class='t-row ng-star-inserted'])[1]//div[contains(@class,'desktop_checkedin')]"
OTHERS_STATUS_XPATH = "xpath=(//div[@class='t-row ng-star-inserted'])[2]//div[contains(@class,'desktop_checkedin')]"

# =========================
# STYLES
# =========================
GREEN_FILL = PatternFill(fill_type="solid", start_color="92D050", end_color="92D050")
HEADER_FILL = PatternFill(fill_type="solid", start_color="D9D9D9", end_color="D9D9D9")
HEADER_FONT = Font(bold=True, size=12)

# =========================
# UTILS
# =========================
def ts():
    return time.strftime("%Y%m%d_%H%M%S")


def today_str():
    return datetime.today().strftime("%Y-%m-%d")


def normalize_str(value):
    return str(value or "").strip()


def backup_corrupt_file(file_path):
    if os.path.exists(file_path):
        new_name = f"{file_path}.{ts()}.corrupt"
        shutil.move(file_path, new_name)
        print(f"Backed up corrupt file: {new_name}")


def backup_file_before_save(file_path):
    if os.path.exists(file_path):
        base = os.path.basename(file_path)
        backup_path = os.path.join(BACKUP_DIR, f"{base}.{ts()}.bak")
        shutil.copy2(file_path, backup_path)
        print(f"Backup created: {backup_path}")


def safe_load_workbook(file_path, create_if_invalid=False):
    if not os.path.exists(file_path):
        return Workbook() if create_if_invalid else None
    try:
        return load_workbook(file_path)
    except Exception as e:
        print(f"Invalid workbook detected: {file_path} | {e}")
        backup_corrupt_file(file_path)
        return Workbook() if create_if_invalid else None


def save_step(page, name):
    try:
        page.screenshot(path=os.path.join(SCREENSHOT_DIR, f"{ts()}_{name}.png"), full_page=True, timeout=8000)
    except Exception:
        pass


# =========================
# EXCEL OUTPUT
# =========================
def sheet_base_name():
    return time.strftime("%d-%m-%y")


def make_unique_sheet_name(output_file):
    base = sheet_base_name()
    if not os.path.exists(output_file):
        return f"{base}(1)"
    wb = safe_load_workbook(output_file, create_if_invalid=True)
    existing = wb.sheetnames
    wb.close()
    i = 1
    while True:
        candidate = f"{base}({i})"
        if candidate not in existing:
            return candidate
        i += 1


def write_results_to_excel(results, output_file):
    if not results:
        return None

    df = pd.DataFrame(results)
    sheet_name = make_unique_sheet_name(output_file)
    wb = safe_load_workbook(output_file, create_if_invalid=True)

    if (
        len(wb.sheetnames) == 1
        and wb.active.max_row == 1
        and wb.active.max_column == 1
        and wb.active["A1"].value is None
    ):
        ws = wb.active
        ws.title = sheet_name
    else:
        ws = wb.create_sheet(title=sheet_name)

    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))

    backup_file_before_save(output_file)
    wb.save(output_file)
    wb.close()
    return sheet_name


# =========================
# TRACKER FUNCTIONS
# =========================
def get_tracker_sheet_name(dt):
    return f"Tracker {dt.strftime('%B')}"


def apply_tracker_header_style(ws):
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def get_header_map(ws):
    return {
        str(ws.cell(row=1, column=i).value).strip(): i
        for i in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=i).value is not None
    }


def ensure_tracker_structure(ws, day_col_name):
    if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
        ws["A1"] = "id"
        ws["B1"] = "email"
        ws["C1"] = "name"
        ws["D1"] = day_col_name
        ws["E1"] = "Total"
        apply_tracker_header_style(ws)
        return

    headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    normalized = [str(h).strip() if h is not None else "" for h in headers]
    required_prefix = ["id", "email", "name"]

    for idx, expected in enumerate(required_prefix, start=1):
        if idx <= len(normalized) and normalized[idx - 1] == expected:
            continue
        if expected not in normalized:
            ws.insert_cols(idx)
            ws.cell(row=1, column=idx).value = expected
            headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
            normalized = [str(h).strip() if h is not None else "" for h in headers]

    header_map = get_header_map(ws)

    if day_col_name not in header_map:
        if "Total" in header_map:
            total_col = header_map["Total"]
            ws.insert_cols(total_col)
            ws.cell(row=1, column=total_col).value = day_col_name
        else:
            ws.cell(row=1, column=ws.max_column + 1).value = day_col_name

    header_map = get_header_map(ws)
    if "Total" not in header_map:
        ws.cell(row=1, column=ws.max_column + 1).value = "Total"

    apply_tracker_header_style(ws)


def find_employee_row(ws, emp_id="", emp_email="", emp_name=""):
    header_map = get_header_map(ws)
    id_col = header_map.get("id")
    email_col = header_map.get("email")
    name_col = header_map.get("name")

    emp_id = normalize_str(emp_id)
    emp_email = normalize_str(emp_email).lower()
    emp_name = normalize_str(emp_name).lower()

    if id_col and emp_id:
        for row_idx in range(2, ws.max_row + 1):
            existing_id = normalize_str(ws.cell(row=row_idx, column=id_col).value)
            if existing_id == emp_id:
                return row_idx

    if email_col and emp_email:
        for row_idx in range(2, ws.max_row + 1):
            existing_email = normalize_str(ws.cell(row=row_idx, column=email_col).value).lower()
            if existing_email == emp_email:
                return row_idx

    if name_col and emp_name:
        for row_idx in range(2, ws.max_row + 1):
            existing_name = normalize_str(ws.cell(row=row_idx, column=name_col).value).lower()
            if existing_name == emp_name:
                return row_idx

    return None


def upsert_employee_row(ws, emp_id="", emp_email="", emp_name=""):
    header_map = get_header_map(ws)
    id_col = header_map["id"]
    email_col = header_map["email"]
    name_col = header_map["name"]

    row_idx = find_employee_row(ws, emp_id, emp_email, emp_name)
    if row_idx is None:
        row_idx = ws.max_row + 1

    current_id = normalize_str(ws.cell(row=row_idx, column=id_col).value)
    current_email = normalize_str(ws.cell(row=row_idx, column=email_col).value)
    current_name = normalize_str(ws.cell(row=row_idx, column=name_col).value)

    if emp_id and not current_id:
        ws.cell(row=row_idx, column=id_col).value = emp_id
    if emp_email and not current_email:
        ws.cell(row=row_idx, column=email_col).value = emp_email
    if emp_name and not current_name:
        ws.cell(row=row_idx, column=name_col).value = emp_name

    return row_idx


def refresh_total_column(ws):
    header_map = get_header_map(ws)
    if "Total" not in header_map:
        return

    total_col = header_map["Total"]
    day_cols = [col_idx for header, col_idx in header_map.items() if header not in ["id", "email", "name", "Total"]]

    for row_idx in range(2, ws.max_row + 1):
        total = 0
        for col_idx in day_cols:
            value = normalize_str(ws.cell(row=row_idx, column=col_idx).value).lower()
            if value == "available":
                total += 1
        ws.cell(row=row_idx, column=total_col).value = total


def update_tracker_excel(results, tracker_file):
    if not results:
        return

    day_col_name = normalize_str(results[0].get("date", today_str()))
    try:
        tracker_date = datetime.strptime(day_col_name, "%Y-%m-%d")
    except Exception:
        tracker_date = datetime.today()

    sheet_name = get_tracker_sheet_name(tracker_date)
    wb = safe_load_workbook(tracker_file, create_if_invalid=True)

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        if (
            len(wb.sheetnames) == 1
            and wb.active.max_row == 1
            and wb.active.max_column == 1
            and wb.active["A1"].value is None
        ):
            ws = wb.active
            ws.title = sheet_name
        else:
            ws = wb.create_sheet(title=sheet_name)

    ensure_tracker_structure(ws, day_col_name)
    header_map = get_header_map(ws)
    day_col = header_map[day_col_name]

    for item in results:
        emp_id = normalize_str(item.get("employee_id", ""))
        emp_email = normalize_str(item.get("employee_email", ""))
        emp_name = normalize_str(item.get("employee_name", ""))
        attendance_value = normalize_str(item.get("attendance", ""))

        if not emp_id and not emp_email and not emp_name:
            continue

        row_idx = upsert_employee_row(ws, emp_id, emp_email, emp_name)
        value = "Available" if attendance_value == "Yes" else "NA"

        cell = ws.cell(row=row_idx, column=day_col)
        cell.value = value
        cell.fill = GREEN_FILL if value == "Available" else PatternFill(fill_type=None)

    refresh_total_column(ws)
    apply_tracker_header_style(ws)
    backup_file_before_save(tracker_file)
    wb.save(tracker_file)
    wb.close()


# =========================
# EMAIL FUNCTIONS
# =========================
def send_html_email(to_email, subject, html_body, plain_text=None, display_name=None, max_retries=MAIL_RETRY_COUNT):
    attempt = 0
    while attempt < max_retries:
        try:
            client = EmailClient.from_connection_string(ACS_CONNECTION)

            recipient = {"address": to_email}
            if display_name:
                recipient["displayName"] = display_name

            message = {
                "senderAddress": ACS_SENDER_EMAIL,
                "recipients": {"to": [recipient]},
                "content": {
                    "subject": subject,
                    "plainText": plain_text or "Your office availability summary is ready.",
                    "html": html_body,
                },
            }

            result = client.begin_send(message).result()
            print(f"Email sent to {to_email}. Message ID: {result['id']}")
            return True

        except Exception as e:
            attempt += 1
            err = str(e)
            if "TooManyRequests" in err:
                wait_time = min(20 * attempt, 120)
                print(f"Rate limited for {to_email}. Retry {attempt}/{max_retries} after {wait_time}s...")
                time.sleep(wait_time)
            else:
                print(f"Failed to send email to {to_email}: {e}")
                return False
    return False


# =========================
# REPORT / AWARENESS
# =========================
def build_awareness_mail_html(employee_name, employee_id, days_in_office, days_pending):
    return f"""
    <html>
    <body style="font-family: Arial, sans-serif;">
        <h2 style="color:#F58025;">Office Availability Summary</h2>
        <p>Hi <b>{employee_name}</b>,</p>
        <p><b>Employee ID:</b> {employee_id}</p>
        <p><b>Days in Office:</b> {days_in_office}</p>
        <p><b>Days Pending:</b> {days_pending}</p>
    </body>
    </html>
    """


def get_month_sheet_or_none(workbook_path):
    wb = safe_load_workbook(workbook_path, create_if_invalid=False)
    if wb is None:
        return None, None, None

    sheet_name = get_tracker_sheet_name(datetime.today())
    if sheet_name not in wb.sheetnames:
        wb.close()
        return None, None, None

    ws = wb[sheet_name]
    return wb, ws, get_header_map(ws)


def clear_failed_awareness_log():
    if os.path.exists(FAILED_AWARENESS_CSV):
        os.remove(FAILED_AWARENESS_CSV)


def read_failed_awareness_recipients():
    if not os.path.exists(FAILED_AWARENESS_CSV):
        return []
    rows = []
    with open(FAILED_AWARENESS_CSV, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows.extend(reader)
    return rows


def log_failed_awareness_mail(employee_id, employee_name, employee_email, reason):
    file_exists = os.path.exists(FAILED_AWARENESS_CSV)
    with open(FAILED_AWARENESS_CSV, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if not file_exists:
            writer.writerow(["date", "employee_id", "employee_name", "employee_email", "reason"])
        writer.writerow([today_str(), employee_id, employee_name, employee_email, reason])


def rewrite_failed_awareness_log(rows):
    if not rows:
        clear_failed_awareness_log()
        return

    with open(FAILED_AWARENESS_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["date", "employee_id", "employee_name", "employee_email", "reason"])
        writer.writerows(rows)


def get_days_in_office_from_test_tracker(employee_id, employee_email="", employee_name=""):
    wb, ws, header_map = get_month_sheet_or_none(TEST_TRACKER_FILE)
    if wb is None:
        return 0

    row_idx = find_employee_row(ws, employee_id, employee_email, employee_name)
    if row_idx is None or "Total" not in header_map:
        wb.close()
        return 0

    try:
        value = int(ws.cell(row=row_idx, column=header_map["Total"]).value or 0)
    except Exception:
        value = 0

    wb.close()
    return value


def send_awareness_mails_from_test_tracker(test_tracker_file):
    wb, ws, header_map = get_month_sheet_or_none(test_tracker_file)
    if wb is None:
        print("Test tracker monthly sheet not found.")
        return

    required = ["id", "email", "name", "Total"]
    for col in required:
        if col not in header_map:
            print(f"Required column missing in test tracker: {col}")
            wb.close()
            return

    clear_failed_awareness_log()
    sent_count = 0
    fail_count = 0

    for row_idx in range(2, ws.max_row + 1):
        employee_id = normalize_str(ws.cell(row=row_idx, column=header_map["id"]).value)
        employee_email = normalize_str(ws.cell(row=row_idx, column=header_map["email"]).value)
        employee_name = normalize_str(ws.cell(row=row_idx, column=header_map["name"]).value)

        try:
            days_in_office = int(ws.cell(row=row_idx, column=header_map["Total"]).value or 0)
        except Exception:
            days_in_office = 0

        if not employee_email:
            continue

        days_pending = max(MONTHLY_TARGET_DAYS - days_in_office, 0)
        html_body = build_awareness_mail_html(employee_name, employee_id, days_in_office, days_pending)

        ok = send_html_email(
            employee_email,
            "Office Availability Summary",
            html_body,
            plain_text=f"Hi {employee_name}, Days in office: {days_in_office}. Days pending: {days_pending}.",
            display_name=employee_name
        )

        if ok:
            sent_count += 1
        else:
            fail_count += 1
            log_failed_awareness_mail(employee_id, employee_name, employee_email, "Failed after retries")

        time.sleep(MAIL_THROTTLE_DELAY_SEC)

    wb.close()
    print(f"Awareness mail run completed. Sent: {sent_count}, Failed: {fail_count}")


def resend_failed_awareness_mails():
    failed_rows = read_failed_awareness_recipients()
    if not failed_rows:
        print("No failed awareness mail log found or no failed recipients.")
        return

    still_failed = []
    success_count = 0

    for row in failed_rows:
        employee_id = normalize_str(row.get("employee_id"))
        employee_name = normalize_str(row.get("employee_name"))
        employee_email = normalize_str(row.get("employee_email"))

        days_in_office = get_days_in_office_from_test_tracker(employee_id, employee_email, employee_name)
        days_pending = max(MONTHLY_TARGET_DAYS - days_in_office, 0)
        html_body = build_awareness_mail_html(employee_name, employee_id, days_in_office, days_pending)

        ok = send_html_email(
            employee_email,
            "Office Availability Summary",
            html_body,
            plain_text=f"Hi {employee_name}, Days in office: {days_in_office}. Days pending: {days_pending}.",
            display_name=employee_name
        )

        if ok:
            success_count += 1
        else:
            still_failed.append([today_str(), employee_id, employee_name, employee_email, "Failed again"])

        time.sleep(MAIL_THROTTLE_DELAY_SEC)

    rewrite_failed_awareness_log(still_failed)
    print(f"Resend summary. Success: {success_count}, Remaining failed: {len(still_failed)}")


def get_working_days_in_month(year, month):
    cal = calendar.monthcalendar(year, month)
    working_days = 0
    for week in cal:
        for day_index in range(0, 5):
            if week[day_index] != 0:
                working_days += 1
    return working_days


def read_tracker_summary(tracker_file):
    wb, ws, header_map = get_month_sheet_or_none(tracker_file)
    if wb is None:
        return None, "Tracker monthly sheet not found."

    if "id" not in header_map or "name" not in header_map or "Total" not in header_map:
        wb.close()
        return None, "Required columns missing."

    rows = []
    total_emp = 0
    total_attendance_sum = 0
    below_10 = 0
    equal_10 = 0
    above_10 = 0

    email_col = header_map.get("email")

    for row_idx in range(2, ws.max_row + 1):
        emp_id = normalize_str(ws.cell(row=row_idx, column=header_map["id"]).value)
        emp_name = normalize_str(ws.cell(row=row_idx, column=header_map["name"]).value)
        emp_email = normalize_str(ws.cell(row=row_idx, column=email_col).value) if email_col else ""

        try:
            total = int(ws.cell(row=row_idx, column=header_map["Total"]).value or 0)
        except Exception:
            total = 0

        if not emp_id and not emp_name and not emp_email:
            continue

        total_emp += 1
        total_attendance_sum += total

        if total < 10:
            below_10 += 1
        elif total == 10:
            equal_10 += 1
        else:
            above_10 += 1

        rows.append({
            "id": emp_id,
            "name": emp_name,
            "email": emp_email,
            "total": total
        })

    wb.close()

    today = datetime.today()
    avg_attendance = round(total_attendance_sum / total_emp, 2) if total_emp else 0
    working_days = get_working_days_in_month(today.year, today.month)

    return {
        "total_emp": total_emp,
        "avg_attendance": avg_attendance,
        "working_days": working_days,
        "below_10": below_10,
        "equal_10": equal_10,
        "above_10": above_10,
        "rows": rows,
        "month_year_text": today.strftime("%B / %Y"),
        "report_date": today.strftime("%Y-%m-%d")
    }, None


def build_month_end_report_html(tracker_file, full_report_url="#", salutation="Team"):
    summary, error = read_tracker_summary(tracker_file)
    if error:
        return f"<html><body><p>{error}</p></body></html>"

    button_html = ""
    if full_report_url:
        button_html = f"""
        <div style="text-align:center; margin-top:20px;">
            <a href="{full_report_url}"
               style="display:inline-block; background:#f58645; color:#ffffff; text-decoration:none; padding:14px 28px; border-radius:8px; font-size:18px; font-weight:500;">
                View Full Report
            </a>
        </div>
        """

    return f"""
    <html>
    <body style="margin:0; padding:0; background:#f3f3f3; font-family:Arial, sans-serif; color:#444;">
        <div style="max-width:800px; margin:20px auto; background:#ffffff; box-shadow:0 2px 10px rgba(0,0,0,0.08);">
            <div style="background:#f58645; padding:36px 40px 30px 40px;">
                <div style="font-size:14px; color:#7a4a2b; letter-spacing:2px; font-weight:bold;">
                    PwC &nbsp;|&nbsp; Canada IT
                </div>

                <h1 style="margin:22px 0 10px 0; font-size:28px; color:#222; font-weight:700;">
                    Monthly Employee Office Availability Report
                </h1>

                <div style="width:72px; height:3px; background:#5b2d1f; margin:8px 0 16px 0;"></div>

                <div style="font-size:16px; color:#5f4638;">
                    {summary['month_year_text']} &nbsp;|&nbsp; Generated on {summary['report_date']}
                </div>
            </div>

            <div style="padding:36px 40px 40px 40px; background:#fafafa;">
                <p style="font-size:16px; margin:0 0 22px 0;">Dear {salutation},</p>

                <p style="font-size:16px; line-height:1.7; margin:0 0 18px 0;">
                    Please find below the consolidated employee office availability report for Canada IT team for
                    <span style="color:#f26a21; font-weight:bold;"> {summary['month_year_text']}</span>.
                    This report presents employee office availability against the monthly benchmark, providing
                    insights into overall alignment with the organization's office presence expectations.
                </p>

                <p style="font-size:16px; line-height:1.7; margin:0 0 30px 0;">
                    Kindly review the summary below. For employee-level details, please click on
                    <b>View Full Report</b>.
                </p>

                <div style="display:flex; gap:12px; margin-bottom:14px; flex-wrap:wrap;">
                    <div style="flex:1; min-width:180px; background:#fff3ea; border:1px solid #f4c9a8; border-radius:8px; text-align:center; padding:18px 10px;">
                        <div style="font-size:22px; font-weight:bold; color:#c65b22;">{summary['total_emp']}</div>
                        <div style="font-size:14px; color:#c97b46;">Total Employees</div>
                    </div>

                    <div style="flex:1; min-width:180px; background:#e9f7ee; border:1px solid #c6e7d1; border-radius:8px; text-align:center; padding:18px 10px;">
                        <div style="font-size:22px; font-weight:bold; color:#1f8b57;">{summary['avg_attendance']}</div>
                        <div style="font-size:14px; color:#339567;">Avg. Days/Employee</div>
                    </div>

                    <div style="flex:1; min-width:180px; background:#fff7ea; border:1px solid #f3ddb4; border-radius:8px; text-align:center; padding:18px 10px;">
                        <div style="font-size:22px; font-weight:bold; color:#bd7a1d;">{summary['working_days']}</div>
                        <div style="font-size:14px; color:#c58a34;">Working Days</div>
                    </div>
                </div>

                <div style="display:flex; gap:12px; margin-bottom:30px; flex-wrap:wrap;">
                    <div style="flex:1; min-width:180px; background:#fde9de; border-radius:10px; text-align:center; padding:24px 10px;">
                        <div style="font-size:22px; font-weight:bold; color:#d45f2e;">{summary['below_10']}</div>
                        <div style="font-size:14px; color:#d45f2e;">Towards Achievement (&lt; 10 days)</div>
                    </div>

                    <div style="flex:1; min-width:180px; background:#dff3e6; border-radius:10px; text-align:center; padding:24px 10px;">
                        <div style="font-size:22px; font-weight:bold; color:#25a05a;">{summary['equal_10']}</div>
                        <div style="font-size:14px; color:#25a05a;">Achieved (= 10 days)</div>
                    </div>

                    <div style="flex:1; min-width:180px; background:#20834d; border-radius:10px; text-align:center; padding:24px 10px;">
                        <div style="font-size:22px; font-weight:bold; color:#ffffff;">{summary['above_10']}</div>
                        <div style="font-size:14px; color:#ffffff;">Above Achievement (&gt; 10 days)</div>
                    </div>
                </div>

                {button_html}
            </div>
        </div>
    </body>
    </html>
    """


def build_full_employee_report_html(tracker_file):
    summary, error = read_tracker_summary(tracker_file)
    if error:
        return f"<html><body><p>{error}</p></body></html>"

    sorted_rows = sorted(summary["rows"], key=lambda x: (-x["total"], x["name"].lower()))

    rows_html = ""
    for idx, item in enumerate(sorted_rows, start=1):
        rows_html += f"""
        <tr>
            <td style="padding:10px; border:1px solid #ddd;">{idx}</td>
            <td style="padding:10px; border:1px solid #ddd;">{item['name']}</td>
            <td style="padding:10px; border:1px solid #ddd;">{item['id']}</td>
            <td style="padding:10px; border:1px solid #ddd;">{item['email']}</td>
            <td style="padding:10px; border:1px solid #ddd; text-align:center;">{item['total']}</td>
        </tr>
        """

    return f"""
    <html>
    <body style="margin:0; padding:30px; background:#f7f7f7; font-family:Arial, sans-serif; color:#333;">
        <div style="max-width:1100px; margin:0 auto; background:#fff; box-shadow:0 2px 10px rgba(0,0,0,0.08);">
            <div style="background:#f58645; padding:30px 35px;">
                <div style="font-size:14px; color:#7a4a2b; letter-spacing:2px; font-weight:bold;">
                    PwC | Canada IT
                </div>
                <h1 style="margin:16px 0 6px 0; color:#222;">Full Employee Office Availability Report</h1>
                <div style="font-size:16px; color:#5f4638;">
                    {summary['month_year_text']} | Generated on {summary['report_date']}
                </div>
            </div>

            <div style="padding:30px 35px;">
                <p style="font-size:16px; margin-bottom:20px;">
                    Detailed employee office availability report.
                </p>

                <table style="width:100%; border-collapse:collapse; background:#fff;">
                    <tr style="background:#f58645; color:#fff;">
                        <th style="padding:12px; border:1px solid #ddd;">#</th>
                        <th style="padding:12px; border:1px solid #ddd;">Employee Name</th>
                        <th style="padding:12px; border:1px solid #ddd;">Employee ID</th>
                        <th style="padding:12px; border:1px solid #ddd;">Employee Email</th>
                        <th style="padding:12px; border:1px solid #ddd;">Total Days in Office</th>
                    </tr>
                    {rows_html}
                </table>
            </div>
        </div>
    </body>
    </html>
    """


def export_month_end_report_html(tracker_file):
    html_body = build_full_employee_report_html(tracker_file)

    dated_export_path = os.path.join(REPORT_EXPORT_DIR, f"month_end_full_report_{today_str()}.html")
    latest_export_path = os.path.join(REPORT_EXPORT_DIR, "month_end_full_report_latest.html")

    with open(dated_export_path, "w", encoding="utf-8") as f:
        f.write(html_body)

    with open(latest_export_path, "w", encoding="utf-8") as f:
        f.write(html_body)

    print(f"Full month-end report exported to archive file: {dated_export_path}")
    print(f"Full month-end report updated at fixed file: {latest_export_path}")
    return latest_export_path


def send_month_end_report_to_abhinandan_and_deborshi(tracker_file):
    export_month_end_report_html(tracker_file)

    recipients = [
        ("Abhinandan Roy", ABHINANDAN_EMAIL),
        ("Deborshi Som", DEBORSHI_EMAIL),
    ]

    for display_name, email in recipients:
        html_body = build_month_end_report_html(
            tracker_file,
            ONEDRIVE_FULL_REPORT_URL,
            salutation=display_name
        )

        plain_text = "Monthly employee office availability report."
        if ONEDRIVE_FULL_REPORT_URL:
            plain_text += f" View full report: {ONEDRIVE_FULL_REPORT_URL}"

        ok = send_html_email(
            email,
            "Monthly Employee Office Availability Report",
            html_body,
            plain_text=plain_text,
            display_name=display_name
        )

        if ok:
            print(f"Month-end report sent successfully to {email}")
        else:
            print(f"Month-end report failed for {email}")

        time.sleep(5)


def send_month_end_report_to_shreyasi(tracker_file):
    export_month_end_report_html(tracker_file)

    html_body = build_month_end_report_html(
        tracker_file,
        ONEDRIVE_FULL_REPORT_URL,
        salutation="Shreyasi Ma'am"
    )

    plain_text = "Monthly employee office availability report."
    if ONEDRIVE_FULL_REPORT_URL:
        plain_text += f" View full report: {ONEDRIVE_FULL_REPORT_URL}"

    ok = send_html_email(
        SHREYASI_EMAIL,
        "Monthly Employee Office Availability Report",
        html_body,
        plain_text=plain_text,
        display_name="Shreyasi Dutta"
    )

    if ok:
        print(f"Month-end report sent successfully to {SHREYASI_EMAIL}")
    else:
        print(f"Month-end report failed for {SHREYASI_EMAIL}")


def prompt_manual_awareness_trigger(test_tracker_file):
    choice = input("Do you want to manually trigger awareness mails to employees from Test_Tracker.xlsx? (yes/no): ").strip().lower()
    if choice == "yes":
        send_awareness_mails_from_test_tracker(test_tracker_file)


def prompt_resend_failed_awareness_trigger():
    choice = input("Do you want to retry failed awareness mails from failed_awareness_mails.csv? (yes/no): ").strip().lower()
    if choice == "yes":
        resend_failed_awareness_mails()


def prompt_export_month_end_report_trigger(tracker_file):
    choice = input("Do you want to export full employee month-end report as HTML? (yes/no): ").strip().lower()
    if choice == "yes":
        export_month_end_report_html(tracker_file)


def prompt_manual_month_end_report_trigger(tracker_file):
    choice = input("Do you want to manually trigger month-end overall report to Abhinandan and Deborshi? (yes/no): ").strip().lower()
    if choice == "yes":
        send_month_end_report_to_abhinandan_and_deborshi(tracker_file)

    choice2 = input("Do you want to manually trigger the same month-end report to Shreyasi ma'am as well? (yes/no): ").strip().lower()
    if choice2 == "yes":
        send_month_end_report_to_shreyasi(tracker_file)


# =========================
# INPUT / MERGE FUNCTIONS
# =========================
def read_employee_input():
    df = pd.read_excel(INPUT_FILE).fillna("")
    rename_map = {}

    for col in df.columns:
        low = str(col).strip().lower()
        if low in ["id", "employee id", "employee_id"]:
            rename_map[col] = "employee_id"
        elif low in ["name", "employee name", "employee_name"]:
            rename_map[col] = "employee_name"
        elif low in ["email", "employee email", "employee_email", "mail"]:
            rename_map[col] = "employee_email"

    df = df.rename(columns=rename_map)

    if "employee_id" not in df.columns:
        df["employee_id"] = ""
    if "employee_name" not in df.columns:
        df["employee_name"] = ""
    if "employee_email" not in df.columns:
        df["employee_email"] = ""

    return df[["employee_id", "employee_name", "employee_email"]]


def read_monthly_employees_from_test_tracker():
    wb, ws, header_map = get_month_sheet_or_none(TEST_TRACKER_FILE)
    if wb is None:
        return pd.DataFrame(columns=["employee_id", "employee_name", "employee_email"])

    rows = []
    for row_idx in range(2, ws.max_row + 1):
        emp_id = normalize_str(ws.cell(row=row_idx, column=header_map["id"]).value) if "id" in header_map else ""
        emp_name = normalize_str(ws.cell(row=row_idx, column=header_map["name"]).value) if "name" in header_map else ""
        emp_email = normalize_str(ws.cell(row=row_idx, column=header_map["email"]).value) if "email" in header_map else ""

        if emp_id or emp_name or emp_email:
            rows.append({
                "employee_id": emp_id,
                "employee_name": emp_name,
                "employee_email": emp_email
            })

    wb.close()
    return pd.DataFrame(rows)


def merge_unique_employees(df1, df2):
    combined = pd.concat(
        [df1[["employee_id", "employee_name", "employee_email"]],
         df2[["employee_id", "employee_name", "employee_email"]]],
        ignore_index=True
    )

    combined["employee_id"] = combined["employee_id"].astype(str).str.strip()
    combined["employee_name"] = combined["employee_name"].astype(str).str.strip()
    combined["employee_email"] = combined["employee_email"].astype(str).str.strip()

    combined = combined[
        (combined["employee_id"] != "") |
        (combined["employee_name"] != "") |
        (combined["employee_email"] != "")
    ].copy()

    combined["merge_key"] = combined.apply(
        lambda r: f"id::{r['employee_id']}" if r["employee_id"]
        else (f"email::{r['employee_email'].lower()}" if r["employee_email"]
        else f"name::{r['employee_name'].lower()}"),
        axis=1
    )

    combined = combined.drop_duplicates(subset=["merge_key"], keep="first").drop(columns=["merge_key"])
    return combined.reset_index(drop=True)


# =========================
# BROWSER REUSE FIX
# =========================
def is_cdp_running(port=CDP_PORT):
    try:
        r = requests.get(f"http://127.0.0.1:{port}/json/version", timeout=2)
        return r.status_code == 200
    except Exception:
        return False


def wait_for_cdp(port=CDP_PORT, retries=30, delay=1):
    for _ in range(retries):
        if is_cdp_running(port):
            return True
        time.sleep(delay)
    return False


def launch_edge_if_needed():
    if is_cdp_running(CDP_PORT):
        print("Reusing existing Edge debug session.")
        return

    print("Launching Edge debug session...")
    subprocess.Popen(
        [
            EDGE_EXE,
            f"--remote-debugging-port={CDP_PORT}",
            f"--user-data-dir={EDGE_PROFILE_DIR}",
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL
    )

    if not wait_for_cdp(CDP_PORT, retries=30, delay=1):
        raise RuntimeError("CDP did not start.")


def get_any_live_page(browser):
    for context in browser.contexts:
        for page in context.pages:
            try:
                if not page.is_closed():
                    return page
            except Exception:
                pass

    context = browser.new_context()
    return context.new_page()


def get_or_open_wis_page(browser):
    for context in browser.contexts:
        for page in context.pages:
            try:
                if not page.is_closed() and "moveinsync.com" in page.url:
                    print("Reusing existing MoveInSync tab.")
                    return page
            except Exception:
                pass

    page = get_any_live_page(browser)
    try:
        page.goto(WIS_URL, wait_until="domcontentloaded")
    except Exception:
        pass
    return page


# =========================
# PLAYWRIGHT HELPERS
# =========================
def text_visible(page, text_value, timeout=5000):
    try:
        page.get_by_text(text_value, exact=False).wait_for(timeout=timeout)
        return True
    except Exception:
        return False


def click_text(page, text_value, step_name):
    save_step(page, f"{step_name}_before")
    page.get_by_text(text_value, exact=False).first.click(timeout=10000)
    time.sleep(4)
    save_step(page, f"{step_name}_after")


def get_wis_frame(page):
    page.locator("iframe").first.wait_for(timeout=20000)
    return page.frame_locator("iframe").first


def open_team_calendar(page):
    frame = get_wis_frame(page)
    try:
        frame.get_by_role("button", name="View Team Calendar").click(timeout=10000, force=True)
        time.sleep(5)
        return True
    except Exception:
        pass

    try:
        frame.locator(VIEW_TEAM_CALENDAR_XPATH).first.click(timeout=10000, force=True)
        time.sleep(5)
        return True
    except Exception:
        pass

    return False


def get_search_box(page):
    frame = get_wis_frame(page)
    try:
        return frame.get_by_role("textbox", name="Search by name, ID or email")
    except Exception:
        return frame.locator(SEARCH_XPATH)


def clear_and_type(search_box, value):
    search_box.click(timeout=5000)
    try:
        search_box.fill("")
        search_box.fill(str(value))
    except Exception:
        search_box.press("Control+A")
        search_box.press("Backspace")
        search_box.type(str(value), delay=100)


def click_clear_all(page):
    frame = get_wis_frame(page)
    try:
        frame.get_by_text("Clear All").click(timeout=5000, force=True)
        time.sleep(1.5)
        return True
    except Exception:
        pass

    try:
        frame.locator(CLEAR_ALL_XPATH).click(timeout=5000, force=True)
        time.sleep(1.5)
        return True
    except Exception:
        return False


def select_employee_from_result(page, emp_id, emp_name, emp_email=""):
    frame = get_wis_frame(page)

    if emp_id:
        try:
            frame.get_by_text(f"ID: {emp_id}", exact=False).first.click(timeout=5000, force=True)
            time.sleep(2)
            return True
        except Exception:
            pass

    if emp_email:
        try:
            frame.get_by_text(emp_email, exact=False).first.click(timeout=5000, force=True)
            time.sleep(2)
            return True
        except Exception:
            pass

    if emp_name:
        try:
            frame.get_by_text(emp_name, exact=False).first.click(timeout=5000, force=True)
            time.sleep(2)
            return True
        except Exception:
            pass

    return False


def close_status_popup_if_any(page):
    frame = get_wis_frame(page)
    try:
        frame.locator("#close-btn").click(timeout=3000, force=True)
        time.sleep(1)
    except Exception:
        pass


def get_container_text(locator, timeout=4000):
    try:
        locator.first.wait_for(timeout=timeout)
        return locator.first.inner_text(timeout=timeout).strip()
    except Exception:
        return ""


def parse_status_text(status_text):
    normalized = " ".join(status_text.upper().split())
    if "CHECKED OUT" in normalized:
        return {"checked_in": "Yes", "checked_out": "Yes"}
    if "CHECKED IN" in normalized:
        return {"checked_in": "Yes", "checked_out": "No"}
    if "BOOKED" in normalized:
        return {"checked_in": "No", "checked_out": "No"}
    return {"checked_in": "No", "checked_out": "No"}


def extract_status_for_abhinandan(page):
    frame = get_wis_frame(page)
    text_val = get_container_text(frame.locator(ABHINANDAN_STATUS_XPATH))
    return parse_status_text(text_val)


def extract_status_for_other_employee(page):
    frame = get_wis_frame(page)
    text_val = get_container_text(frame.locator(OTHERS_STATUS_XPATH))
    return parse_status_text(text_val)


def get_attendance_value(checked_in, checked_out):
    return "Yes" if checked_in == "Yes" or checked_out == "Yes" else "No"


# =========================
# MAIN
# =========================
def main():
    launch_edge_if_needed()

    input_employees = read_employee_input()
    test_tracker_employees = read_monthly_employees_from_test_tracker()
    employees = merge_unique_employees(input_employees, test_tracker_employees)

    print(f"Total unique employees to process: {len(employees)}")
    results = []

    with sync_playwright() as p:
        browser = p.chromium.connect_over_cdp(f"http://127.0.0.1:{CDP_PORT}")
        page = get_or_open_wis_page(browser)
        page.set_default_timeout(TIMEOUT_MS)

        try:
            time.sleep(5)

            if text_visible(page, "Pick an account", 5000):
                try:
                    page.locator('[data-test-id="abhinandan.roy@pwc.com"]').click(timeout=10000, force=True)
                    time.sleep(8)
                except Exception:
                    try:
                        click_text(page, "abhinandan.roy@pwc.com", "pick_account")
                    except Exception:
                        pass

            opened = open_team_calendar(page)
            if not opened:
                input("Open Team Calendar manually, then press Enter... ")

            for _, row in employees.iterrows():
                emp_id = normalize_str(row.get("employee_id", ""))
                emp_name = normalize_str(row.get("employee_name", ""))
                emp_email = normalize_str(row.get("employee_email", ""))

                if not emp_id and not emp_name and not emp_email:
                    continue

                try:
                    if emp_id == ABHINANDAN_ID or emp_name.lower() == ABHINANDAN_NAME.lower():
                        status = extract_status_for_abhinandan(page)
                    else:
                        click_clear_all(page)
                        search_box = get_search_box(page)
                        search_value = emp_id if emp_id else (emp_email if emp_email else emp_name)
                        clear_and_type(search_box, search_value)
                        time.sleep(SEARCH_WAIT_SEC)

                        selected = select_employee_from_result(page, emp_id, emp_name, emp_email)
                        if not selected:
                            results.append({
                                "date": today_str(),
                                "employee_id": emp_id,
                                "employee_name": emp_name,
                                "employee_email": emp_email,
                                "attendance": "No",
                                "checked_in": "No",
                                "checked_out": "No",
                            })
                            print(f"Not found: {emp_name} | {emp_id} | {emp_email}")
                            continue

                        status = extract_status_for_other_employee(page)
                        close_status_popup_if_any(page)

                    attendance = get_attendance_value(status["checked_in"], status["checked_out"])

                    results.append({
                        "date": today_str(),
                        "employee_id": emp_id,
                        "employee_name": emp_name,
                        "employee_email": emp_email,
                        "attendance": attendance,
                        "checked_in": status["checked_in"],
                        "checked_out": status["checked_out"],
                    })

                    print(f"Processed: {emp_name} | {emp_id} | {emp_email} | {attendance}")

                except Exception as e:
                    print(f"Error for {emp_name} | {emp_id} | {emp_email}: {e}")
                    results.append({
                        "date": today_str(),
                        "employee_id": emp_id,
                        "employee_name": emp_name,
                        "employee_email": emp_email,
                        "attendance": "Error",
                        "checked_in": "Error",
                        "checked_out": "Error",
                    })

            sheet_name = write_results_to_excel(results, OUTPUT_FILE)
            update_tracker_excel(results, TRACKER_FILE)
            update_tracker_excel(results, TEST_TRACKER_FILE)

            print(f"Done. Output saved to {OUTPUT_FILE}, sheet: {sheet_name}")
            print(f"Updated {TRACKER_FILE}")
            print(f"Updated {TEST_TRACKER_FILE}")

            prompt_manual_awareness_trigger(TEST_TRACKER_FILE)
            prompt_resend_failed_awareness_trigger()
            prompt_export_month_end_report_trigger(TRACKER_FILE)
            prompt_manual_month_end_report_trigger(TRACKER_FILE)

        finally:
            browser.close()


if __name__ == "__main__":
    main()